"""Ingest Google Drive folder downloads (zip archives and loose files) into Qdrant."""

from __future__ import annotations

import io
import re
import tempfile
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

from rag.nrag.data import Chunk, Document, Paragraph, TextParagraph
from rag.nrag.ids import stable_point_id
from rag.nrag.pipeline import upsert_documents
from rag.nrag.split import (
    DEFAULT_CHUNK_CHARS,
    DEFAULT_CHUNK_OVERLAP,
    LengthSplitStrategy,
    normalize_whitespace,
    split_text,
)

# Google Drive "Download" of a folder usually yields one or more zips of exported files.
_TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".log",
    ".html",
    ".htm",
}
_DOCX_SUFFIXES = {".docx"}
_PDF_SUFFIXES = {".pdf"}
_XLSX_SUFFIXES = {".xlsx", ".xlsm"}
_MARKDOWN_SUFFIXES = {".md", ".markdown"}
_SKIP_DIR_NAMES = {"__macosx", ".git", ".svn", "node_modules"}
_SKIP_FILE_NAMES = {".ds_store", "thumbs.db", "desktop.ini"}
_SKIP_SUFFIXES = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".svg",
    ".ico",
    ".mp4",
    ".mov",
    ".zip",
    ".gz",
    ".tgz",
    ".7z",
    ".exe",
    ".dmg",
    ".gdoc",
    ".gsheet",
    ".gslides",
    ".gdraw",
    ".gform",
    ".gtable",
}

# Re-export for tests and callers that imported these from this module.
chunk_text = split_text


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
        if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"}:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = data.strip()
        if text:
            self._chunks.append(text)

    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", " ".join(self._chunks)).strip()


def _read_text_bytes(data: bytes, path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            raw = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raw = data.decode("utf-8", errors="ignore")
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(raw)
        return normalize_whitespace(parser.text())
    return normalize_whitespace(raw)


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return normalize_whitespace("\n".join(parts))


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(text)
    return normalize_whitespace("\n\n".join(parts))


def _extract_xlsx(data: bytes) -> str:
    import pandas as pd

    frames = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str, engine="openpyxl")
    parts: list[str] = []
    for sheet_name, frame in frames.items():
        frame = frame.fillna("")
        parts.append(f"# Sheet: {sheet_name}")
        preview = frame.head(200)
        for row in preview.itertuples(index=False, name=None):
            cells = [str(cell).strip() for cell in row if str(cell).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return normalize_whitespace("\n".join(parts))


def extract_text_from_bytes(data: bytes, path: Path) -> str | None:
    """Return plain text for a Drive-exported file, or None if unsupported/empty."""
    suffix = path.suffix.lower()
    name = path.name.lower()
    if name in _SKIP_FILE_NAMES or suffix in _SKIP_SUFFIXES:
        return None
    try:
        if suffix in _TEXT_SUFFIXES:
            text = _read_text_bytes(data, path)
        elif suffix in _DOCX_SUFFIXES:
            text = _extract_docx(data)
        elif suffix in _PDF_SUFFIXES:
            text = _extract_pdf(data)
        elif suffix in _XLSX_SUFFIXES:
            text = _extract_xlsx(data)
        else:
            return None
    except Exception as exc:  # noqa: BLE001 — skip bad files, continue ingest
        print(f"Skipping {path}: extract failed ({exc})")
        return None
    return text or None


@dataclass(kw_only=True)
class DriveDocument(Document):
    source_zip: str = ""
    rel_path: str = ""
    suffix: str = ""

    def split_paragraphs(self) -> list[Paragraph]:
        if self.paragraphs:
            return self.paragraphs
        if self.suffix in _MARKDOWN_SUFFIXES:
            from rag.nrag.pdf import markdown_section_texts

            paras: list[Paragraph] = []
            for seq, (title, body) in enumerate(
                markdown_section_texts(self.content, fallback_title=self.title)
            ):
                paras.append(
                    TextParagraph(
                        ref_doc=self,
                        para_seq=seq,
                        para_title=title or self.title,
                        content=body,
                    )
                )
            if paras:
                self.paragraphs = paras
                return self.paragraphs
        return super().split_paragraphs()

    def _file_chunk_index(self, chunk: Chunk) -> int:
        index = chunk.chunk_seq
        for para in self.paragraphs:
            if para is chunk.ref_para:
                return index
            index += len(para.chunks)
        return index

    def chunk_payload(self, chunk: Chunk) -> dict[str, object]:
        return {
            "source": "google_drive",
            "source_zip": self.source_zip,
            "path": self.rel_path,
            "title": self.title,
            "suffix": self.suffix,
            "chunk_index": self._file_chunk_index(chunk),
            "text": chunk.content,
        }

    def chunk_embed_text(self, chunk: Chunk) -> str:
        return f"Title: {self.title}\nPath: {self.rel_path}\n\n{chunk.content}"

    def chunk_point_id(self, chunk: Chunk) -> int | str:
        # Single-paragraph files keep the historical ``prefix::path::index`` key.
        if len(self.paragraphs) <= 1 and chunk.ref_para.para_seq == 0:
            key = f"{self.doc_id}::{self.rel_path}::{chunk.chunk_seq}"
        else:
            key = (
                f"{self.doc_id}::{self.rel_path}"
                f"::{chunk.ref_para.para_seq}::{chunk.chunk_seq}"
            )
        return stable_point_id(key)


def _iter_files(root: Path, *, include_zip: bool = False) -> Iterable[Path]:
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if any(part.lower() in _SKIP_DIR_NAMES for part in path.parts):
            continue
        if path.name.lower() in _SKIP_FILE_NAMES:
            continue
        if not include_zip and path.suffix.lower() == ".zip":
            continue
        yield path


def _documents_from_tree(
    root: Path,
    *,
    source_zip: str,
    relative_to: Path,
    point_prefix: str,
) -> list[DriveDocument]:
    docs: list[DriveDocument] = []
    for path in _iter_files(root):
        try:
            data = path.read_bytes()
        except OSError as exc:
            print(f"Skipping {path}: read failed ({exc})")
            continue
        text = extract_text_from_bytes(data, path)
        if not text:
            continue
        try:
            rel = str(path.relative_to(relative_to))
        except ValueError:
            rel = path.name
        docs.append(
            DriveDocument(
                doc_id=point_prefix,
                url=rel,
                title=path.stem,
                content=text,
                source_zip=source_zip,
                rel_path=rel,
                suffix=path.suffix.lower(),
            )
        )
    return docs


def collect_google_drive_documents(directory: str | Path) -> list[DriveDocument]:
    """Scan ``directory`` for Drive zips and loose files; return documents to split."""
    root = Path(directory).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Google Drive ingest directory not found: {root}")

    docs: list[DriveDocument] = []
    zips = sorted(p for p in root.iterdir() if p.is_file() and p.suffix.lower() == ".zip")
    for zip_path in zips:
        print(f"Extracting {zip_path.name} ...")
        with tempfile.TemporaryDirectory(prefix="gdrive_rag_") as tmp:
            tmp_root = Path(tmp)
            try:
                with zipfile.ZipFile(zip_path, "r") as archive:
                    archive.extractall(tmp_root)
            except zipfile.BadZipFile as exc:
                print(f"Skipping {zip_path.name}: bad zip ({exc})")
                continue
            docs.extend(
                _documents_from_tree(
                    tmp_root,
                    source_zip=zip_path.name,
                    relative_to=tmp_root,
                    point_prefix=zip_path.name,
                )
            )

    docs.extend(
        _documents_from_tree(
            root,
            source_zip="",
            relative_to=root,
            point_prefix="loose",
        )
    )

    unique = {f"{doc.doc_id}::{doc.rel_path}": doc for doc in docs}
    return list(unique.values())


def upload_google_drive(
    directory: str | Path,
    collection_name: str,
    *,
    batch_size: int | None = None,
) -> int:
    """Embed Google Drive downloads under ``directory`` into ``collection_name``.

    ``directory`` should contain the zip file(s) downloaded from Google Drive
    (and optionally already-extracted files). Returns the number of upserted points.
    """
    docs = collect_google_drive_documents(directory)
    print(f"Collected {len(docs)} files from {directory!s} → collection {collection_name!r}")
    return upsert_documents(
        collection_name,
        docs,
        strategy=LengthSplitStrategy(
            chunk_chars=DEFAULT_CHUNK_CHARS,
            overlap=DEFAULT_CHUNK_OVERLAP,
        ),
        batch_size=batch_size,
    )
