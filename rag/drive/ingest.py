"""Ingest extracted Google Drive folder downloads into Qdrant.

Zip archives are skipped — extract them first, then point ingest at the tree.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

from rag.document import Chunk, Document, Paragraph, TextParagraph, stable_point_id
from rag.ingest import length_strategy_for_embeddings, upsert_documents
from rag.chunking import (
    normalize_whitespace,
    split_text,
)

# Extracted Drive folder downloads (zip files in the tree are skipped).
_TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".log",
    ".html",
    ".htm",
}
_DOCX_SUFFIXES = {".docx"}
_PDF_SUFFIXES = {".pdf"}
_XLSX_SUFFIXES = {".xlsx", ".xlsm"}
_MARKDOWN_SUFFIXES = {".md", ".markdown"}
_SKIP_DIR_NAMES = {"__macosx", ".git", ".svn", "node_modules"}
_SKIP_FILE_NAMES = {".ds_store", "thumbs.db", "desktop.ini"}
_SKIP_SUFFIXES = {
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".svg",
    ".ico",
    ".mp4",
    ".mov",
    ".zip",
    ".gz",
    ".tgz",
    ".7z",
    ".exe",
    ".dmg",
    ".gdoc",
    ".gsheet",
    ".gslides",
    ".gdraw",
    ".gform",
    ".gtable",
}

# Re-export for tests and callers that imported these from this module.
chunk_text = split_text


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
        if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"}:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = data.strip()
        if text:
            self._chunks.append(text)

    def text(self) -> str:
        return re.sub(r"\n{3,}", "\n\n", " ".join(self._chunks)).strip()


def _read_text_bytes(data: bytes, path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            raw = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raw = data.decode("utf-8", errors="ignore")
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(raw)
        return normalize_whitespace(parser.text())
    return normalize_whitespace(raw)


def _extract_docx(data: bytes) -> str:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return normalize_whitespace("\n".join(parts))


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(text)
    return normalize_whitespace("\n\n".join(parts))


def _extract_xlsx(data: bytes) -> str:
    import pandas as pd

    frames = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str, engine="openpyxl")
    parts: list[str] = []
    for sheet_name, frame in frames.items():
        frame = frame.fillna("")
        parts.append(f"# Sheet: {sheet_name}")
        preview = frame.head(200)
        for row in preview.itertuples(index=False, name=None):
            cells = [str(cell).strip() for cell in row if str(cell).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return normalize_whitespace("\n".join(parts))


def extract_text_from_bytes(data: bytes, path: Path) -> str | None:
    """Return plain text for a Drive-exported file, or None if unsupported/empty."""
    suffix = path.suffix.lower()
    name = path.name.lower()
    if name in _SKIP_FILE_NAMES or suffix in _SKIP_SUFFIXES:
        return None
    try:
        if suffix in _TEXT_SUFFIXES:
            text = _read_text_bytes(data, path)
        elif suffix in _DOCX_SUFFIXES:
            text = _extract_docx(data)
        elif suffix in _PDF_SUFFIXES:
            text = _extract_pdf(data)
        elif suffix in _XLSX_SUFFIXES:
            text = _extract_xlsx(data)
        else:
            return None
    except Exception as exc:  # noqa: BLE001 — skip bad files, continue ingest
        print(f"Skipping {path}: extract failed ({exc})")
        return None
    return text or None


@dataclass(kw_only=True)
class DriveDocument(Document):
    source_zip: str = ""
    rel_path: str = ""
    suffix: str = ""

    def split_paragraphs(self) -> list[Paragraph]:
        if self.paragraphs:
            return self.paragraphs
        if self.suffix in _MARKDOWN_SUFFIXES:
            from rag.markdown import markdown_section_texts

            paras: list[Paragraph] = []
            for seq, (title, body) in enumerate(
                markdown_section_texts(self.content, fallback_title=self.title)
            ):
                paras.append(
                    TextParagraph(
                        ref_doc=self,
                        para_seq=seq,
                        para_title=title or self.title,
                        content=body,
                    )
                )
            if paras:
                self.paragraphs = paras
                return self.paragraphs
        return super().split_paragraphs()

    def _file_chunk_index(self, chunk: Chunk) -> int:
        index = chunk.chunk_seq
        for para in self.paragraphs:
            if para is chunk.ref_para:
                return index
            index += len(para.chunks)
        return index

    def chunk_payload(self, chunk: Chunk) -> dict[str, object]:
        return {
            "source": "google_drive",
            "source_zip": self.source_zip,
            "path": self.rel_path,
            "title": self.title,
            "suffix": self.suffix,
            "chunk_index": self._file_chunk_index(chunk),
            "text": chunk.content,
        }

    def chunk_embed_text(self, chunk: Chunk) -> str:
        return f"Title: {self.title}\nPath: {self.rel_path}\n\n{chunk.content}"

    def chunk_point_id(self, chunk: Chunk) -> int | str:
        # Single-paragraph files keep the historical ``prefix::path::index`` key.
        if len(self.paragraphs) <= 1 and chunk.ref_para.para_seq == 0:
            key = f"{self.doc_id}::{self.rel_path}::{chunk.chunk_seq}"
        else:
            key = (
                f"{self.doc_id}::{self.rel_path}"
                f"::{chunk.ref_para.para_seq}::{chunk.chunk_seq}"
            )
        return stable_point_id(key)


def _iter_files(root: Path) -> Iterable[Path]:
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if any(part.lower() in _SKIP_DIR_NAMES for part in path.parts):
            continue
        if path.name.lower() in _SKIP_FILE_NAMES:
            continue
        if path.suffix.lower() == ".zip":
            continue
        yield path


def list_drive_paths(directory: str | Path) -> list[Path]:
    """Return extracted files under ``directory``. Zip archives are skipped."""
    root = Path(directory).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"Google Drive ingest directory not found: {root}")
    zips = sorted(p for p in root.rglob("*.zip") if p.is_file())
    if zips:
        print(
            f"Skipping {len(zips)} zip file(s) under {root}; "
            "extract them first, then re-run plan."
        )
    return sorted(_iter_files(root))


def documents_from_path(
    path: str | Path,
    *,
    ingest_root: str | Path,
) -> list[DriveDocument]:
    """Load documents from one extracted Drive file. Zip paths are skipped."""
    file_path = Path(path).expanduser().resolve()
    root = Path(ingest_root).expanduser().resolve()
    if file_path.suffix.lower() == ".zip":
        print(f"Skipping {file_path.name}: zip files are not ingested; extract first")
        return []
    try:
        data = file_path.read_bytes()
    except OSError as exc:
        raise OSError(f"Failed to read {file_path}: {exc}") from exc
    text = extract_text_from_bytes(data, file_path)
    if not text:
        return []
    try:
        rel = str(file_path.relative_to(root))
    except ValueError:
        rel = file_path.name
    return [
        DriveDocument(
            doc_id="loose",
            url=rel,
            title=file_path.stem,
            content=text,
            source_zip="",
            rel_path=rel,
            suffix=file_path.suffix.lower(),
        )
    ]


def collect_google_drive_documents(directory: str | Path) -> list[DriveDocument]:
    """Scan ``directory`` for extracted Drive files; zip archives are skipped."""
    root = Path(directory).expanduser().resolve()
    docs: list[DriveDocument] = []
    for path in list_drive_paths(root):
        docs.extend(documents_from_path(path, ingest_root=root))
    unique = {doc.rel_path: doc for doc in docs}
    return list(unique.values())


def upload_google_drive(
    directory: str | Path,
    collection_name: str,
    *,
    batch_size: int | None = None,
) -> int:
    """Embed extracted Google Drive files under ``directory`` into ``collection_name``.

    ``directory`` should already contain unzipped files. Zip archives are skipped.
    Returns the number of upserted points.
    """
    docs = collect_google_drive_documents(directory)
    print(f"Collected {len(docs)} files from {directory!s} → collection {collection_name!r}")
    return upsert_documents(
        collection_name,
        docs,
        strategy=length_strategy_for_embeddings(),
        batch_size=batch_size,
    )
